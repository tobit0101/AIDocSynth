from aidocsynth.services.metadata_service import MetadataService
import fitz
from PIL import Image
from docx import Document
from pptx import Presentation
import openpyxl


def test_generate_and_merge_metadata_basic_merge():
    svc = MetadataService()
    classification_data = {
        "headline": "New Title",
        "author": "Alice",
        "keywords": ["foo", "bar"],
        "description": "Short desc",
    }
    original = {
        "title": "Old Title",
        "subject": "Old Subject",
    }

    merged = svc.generate_and_merge_metadata(classification_data, original)
    assert merged["title"] == "New Title"  # headline -> title
    assert merged["author"] == "Alice"
    assert merged["keywords"] == "foo, bar"
    assert merged["description"] == "Short desc"
    # Unspecified mappings remain
    assert merged["subject"] == "Old Subject"


def test_pdf_metadata_roundtrip(tmp_path):
    svc = MetadataService()
    pdf_path = tmp_path / "meta.pdf"

    # Create a minimal PDF
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    # Write metadata
    meta = {
        "title": "T1",
        "author": "A1",
        "subject": "S1",
        "keywords": "k1, k2",
        "creator_tool": "UnitTest",
    }
    assert svc.set_file_metadata(pdf_path, meta) is True

    # Read back and verify selected fields
    got = svc.get_file_metadata(pdf_path)
    assert got.get("title") == "T1"
    assert got.get("author") == "A1"
    assert got.get("subject") == "S1"
    assert got.get("keywords") == "k1, k2"
    assert got.get("creator_tool") == "UnitTest"


def test_pdf_set_metadata_no_updates_returns_true(tmp_path):
    svc = MetadataService()
    pdf_path = tmp_path / "empty.pdf"

    # Minimal PDF
    doc = fitz.open()
    doc.new_page()
    doc.save(str(pdf_path))
    doc.close()

    # No known keys -> update_dict is empty, should still return True
    assert svc.set_file_metadata(pdf_path, {}) is True


def test_png_metadata_roundtrip(tmp_path):
    svc = MetadataService()
    png_path = tmp_path / "img.png"

    # Create a tiny PNG
    Image.new("RGB", (2, 2), color=(255, 255, 255)).save(png_path)

    meta = {
        "author": "Alice",
        "title": "Hello",
        "description": "World",
    }
    assert svc.set_file_metadata(png_path, meta) is True

    got = svc.get_file_metadata(png_path)
    # PNG handler maps textual fields directly
    assert got.get("author") == "Alice"
    assert got.get("title") == "Hello"
    assert got.get("description") == "World"


def test_unsupported_extension_returns_defaults(tmp_path):
    svc = MetadataService()
    p = tmp_path / "note.txt"
    p.write_text("hello")

    # Reading unsupported -> {}
    assert svc.get_file_metadata(p) == {}
    # Writing unsupported -> False
    assert svc.set_file_metadata(p, {"title": "x"}) is False


def test_docx_metadata_roundtrip(tmp_path):
    svc = MetadataService()
    path = tmp_path / "meta.docx"

    # Create minimal DOCX
    Document().save(str(path))

    meta = {
        "title": "DocTitle",
        "author": "DocAuthor",
        "subject": "DocSubject",
        "keywords": "k1, k2",
        "last_modified_by": "Tester",
        "description": "DocComments",
    }
    assert svc.set_file_metadata(path, meta) is True

    got = svc.get_file_metadata(path)
    assert got.get("title") == "DocTitle"
    assert got.get("author") == "DocAuthor"
    assert got.get("subject") == "DocSubject"
    assert got.get("keywords") == "k1, k2"
    # comments mapped to description
    assert got.get("description") == "DocComments"


def test_docx_invalid_file_returns_empty_and_false(tmp_path):
    svc = MetadataService()
    bad = tmp_path / "bad.docx"
    bad.write_text("not a real docx")

    assert svc.get_file_metadata(bad) == {}
    assert svc.set_file_metadata(bad, {"title": "x"}) is False


def test_xlsx_metadata_roundtrip(tmp_path):
    svc = MetadataService()
    path = tmp_path / "meta.xlsx"

    # Create minimal XLSX
    wb = openpyxl.Workbook()
    wb.save(str(path))

    meta = {
        "title": "XTitle",
        "author": "XAuthor",
        "subject": "XSubject",
        "keywords": "x1, x2",
        "last_modified_by": "XUser",
        "description": "XDesc",
    }
    assert svc.set_file_metadata(path, meta) is True

    got = svc.get_file_metadata(path)
    assert got.get("title") == "XTitle"
    assert got.get("author") == "XAuthor"
    assert got.get("subject") == "XSubject"
    assert got.get("keywords") == "x1, x2"
    assert got.get("last_modified_by") == "XUser"
    assert got.get("description") == "XDesc"


def test_pptx_metadata_roundtrip(tmp_path):
    svc = MetadataService()
    path = tmp_path / "meta.pptx"

    # Create minimal PPTX
    Presentation().save(str(path))

    meta = {
        "title": "PTit",
        "author": "PAutor",
        "subject": "PSub",
        "keywords": "p1, p2",
        "last_modified_by": "PUser",
        "description": "PDesc",
    }
    assert svc.set_file_metadata(path, meta) is True

    got = svc.get_file_metadata(path)
    assert got.get("title") == "PTit"
    assert got.get("author") == "PAutor"
    assert got.get("subject") == "PSub"
    assert got.get("keywords") == "p1, p2"
    assert got.get("last_modified_by") == "PUser"
    assert got.get("description") == "PDesc"


def test_jpeg_exif_metadata_roundtrip(tmp_path):
    svc = MetadataService()
    jpg_path = tmp_path / "img.jpg"

    # Create a tiny JPEG
    Image.new("RGB", (4, 4), color=(200, 100, 50)).save(jpg_path, format="JPEG")

    meta = {
        "author": "Alice",
        "description": "Desc",
        "comment": "Note",
    }
    assert svc.set_file_metadata(jpg_path, meta) is True

    got = svc.get_file_metadata(jpg_path)
    # EXIF handler maps author/description/comment from proper tags
    assert got.get("author") == "Alice"
    assert got.get("description") == "Desc"
    assert got.get("comment") == "Note"


def test_safe_decode_branches():
    svc = MetadataService()
    assert svc._safe_decode(None) == ""
    assert svc._safe_decode(b"abc") == "abc"
    assert svc._safe_decode("hello") == "hello"


def test_png_get_ignores_unmapped_text_keys(tmp_path):
    from PIL import PngImagePlugin

    svc = MetadataService()
    p = tmp_path / "text.png"
    img = Image.new("RGB", (2, 2), color=(10, 10, 10))
    info = PngImagePlugin.PngInfo()
    info.add_text("Software", "Pillow")
    info.add_text("RandomKey", "Value")
    img.save(p, pnginfo=info)

    got = svc.get_file_metadata(p)
    # Unmapped keys should not appear
    assert got == {}


def test_jpeg_get_exif_only_comment_branch(tmp_path):
    import piexif

    svc = MetadataService()
    p = tmp_path / "exif_only.jpg"
    # Create EXIF with only the UserComment in Exif IFD, 0th empty
    exif_dict = {"0th": {}, "Exif": {piexif.ExifIFD.UserComment: b"Cmt"}, "GPS": {}, "1st": {}, "thumbnail": None}
    exif_bytes = piexif.dump(exif_dict)
    Image.new("RGB", (3, 3), color=(50, 60, 70)).save(p, format="JPEG", exif=exif_bytes)

    got = svc.get_file_metadata(p)
    assert got.get("comment") == "Cmt"


def test_jpeg_get_exif_only_author_description_branch(tmp_path):
    import piexif

    svc = MetadataService()
    p = tmp_path / "exif_0th_only.jpg"
    # Only 0th IFD tags present
    exif_dict = {
        "0th": {
            piexif.ImageIFD.Artist: b"Alice",
            piexif.ImageIFD.ImageDescription: b"Desc",
        },
        "Exif": {}, "GPS": {}, "1st": {}, "thumbnail": None,
    }
    exif_bytes = piexif.dump(exif_dict)
    Image.new("RGB", (3, 3), color=(90, 100, 110)).save(p, format="JPEG", exif=exif_bytes)

    got = svc.get_file_metadata(p)
    assert got.get("author") == "Alice"
    assert got.get("description") == "Desc"


def test_get_file_metadata_generic_exception_returns_empty(tmp_path, caplog):
    svc = MetadataService()
    p = tmp_path / "img.png"
    Image.new("RGB", (2, 2), color=(1, 2, 3)).save(p)

    def boom(_):
        raise RuntimeError("boom")

    orig_get, orig_set = svc._metadata_handlers[".png"]
    svc._metadata_handlers[".png"] = (boom, orig_set)

    with caplog.at_level("ERROR"):
        out = svc.get_file_metadata(p)
    assert out == {}


def test_set_file_metadata_generic_exception_returns_false(tmp_path, caplog):
    svc = MetadataService()
    p = tmp_path / "img2.png"
    Image.new("RGB", (2, 2), color=(1, 2, 3)).save(p)

    def boom(_, __):
        raise RuntimeError("boom")

    orig_get, orig_set = svc._metadata_handlers[".png"]
    svc._metadata_handlers[".png"] = (orig_get, boom)

    with caplog.at_level("ERROR"):
        ok = svc.set_file_metadata(p, {"title": "T"})
    assert ok is False


def test_get_file_metadata_specific_exceptions_pptx_xlsx(tmp_path):
    svc = MetadataService()

    bad_pptx = tmp_path / "bad.pptx"
    bad_pptx.write_text("not a real pptx")
    assert svc.get_file_metadata(bad_pptx) == {}

    bad_xlsx = tmp_path / "bad.xlsx"
    bad_xlsx.write_text("not a real xlsx")
    assert svc.get_file_metadata(bad_xlsx) == {}


def test_set_file_metadata_specific_exceptions_pptx_xlsx(tmp_path):
    svc = MetadataService()

    bad_pptx = tmp_path / "bad_set.pptx"
    bad_pptx.write_text("invalid")
    assert svc.set_file_metadata(bad_pptx, {"title": "X"}) is False

    bad_xlsx = tmp_path / "bad_set.xlsx"
    bad_xlsx.write_text("invalid")
    assert svc.set_file_metadata(bad_xlsx, {"title": "X"}) is False


def test_docx_set_noop_path(tmp_path):
    svc = MetadataService()
    d = tmp_path / "noop.docx"
    Document().save(str(d))

    # Capture initial metadata (python-docx sets some defaults)
    before = svc.get_file_metadata(d)
    # No keys -> modified remains False, should return True and not change metadata
    assert svc.set_file_metadata(d, {}) is True
    after = svc.get_file_metadata(d)
    assert after == before


def test_xlsx_set_noop_path(tmp_path):
    svc = MetadataService()
    path = tmp_path / "noop.xlsx"
    wb = openpyxl.Workbook()
    wb.save(str(path))

    before = svc.get_file_metadata(path)
    assert svc.set_file_metadata(path, {}) is True
    after = svc.get_file_metadata(path)
    assert after == before


def test_pptx_set_noop_path(tmp_path):
    svc = MetadataService()
    path = tmp_path / "noop.pptx"
    Presentation().save(str(path))

    before = svc.get_file_metadata(path)
    assert svc.set_file_metadata(path, {}) is True
    after = svc.get_file_metadata(path)
    assert after == before
