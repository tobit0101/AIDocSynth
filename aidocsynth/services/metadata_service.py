import logging
from pathlib import Path
from typing import Dict, Any, Optional, Callable, Tuple, List

# Imports for metadata handling
import fitz  # PyMuPDF
from PIL import Image, PngImagePlugin
import piexif
from docx import Document
from docx.opc.exceptions import PackageNotFoundError as DocxPackageNotFoundError
from pptx import Presentation
from pptx.exc import PackageNotFoundError as PptxPackageNotFoundError
import openpyxl
from openpyxl.utils.exceptions import InvalidFileException as OpenpyxlInvalidFileException

logger = logging.getLogger(__name__)

COMMON_METADATA_KEYS = {
    "title": "title",
    "author": "author",
    "subject": "subject",
    "keywords": "keywords",
    "creator_tool": "creator_tool",
    "creation_date": "creation_date",
    "modification_date": "modification_date",
    "last_modified_by": "last_modified_by",
    "description": "description",
    "comment": "comment"
}


def write_metadata_task(new_path, classification_data, original_metadata):
    """
    A picklable top-level function to be executed in a separate process.
    It generates, merges, and writes metadata to a file.
    """
    # Since this is a top-level function, we instantiate the service here.
    metadata_service = MetadataService()
    final_metadata = metadata_service.generate_and_merge_metadata(
        classification_data,
        original_metadata
    )
    if final_metadata:
        metadata_service.set_file_metadata(new_path, final_metadata)


class MetadataService:
    """Handles reading, writing, and merging of file metadata."""

    def __init__(self):
        self._metadata_handlers = self._register_metadata_handlers()

    def generate_and_merge_metadata(
        self, 
        classification_data: Dict[str, Any],
        original_metadata: Dict[str, Any]
    ) -> Dict[str, Any]:
        """
        Merges LLM-generated metadata with original file metadata.
        """
        final_metadata = original_metadata.copy()
        llm_mapping = {
            'headline': 'title',
            'context': 'subject',
            'author': 'author',
            'comment': 'comment',
            'keywords': 'keywords'
        }

        for llm_key, meta_key in llm_mapping.items():
            if llm_key in classification_data and classification_data[llm_key]:
                value = classification_data[llm_key]
                if isinstance(value, list):
                    final_metadata[meta_key] = ', '.join(value)
                else:
                    final_metadata[meta_key] = str(value)
        
        logger.debug(f"Final merged metadata: {final_metadata}")
        return final_metadata

    # --- Metadata Strategy --- #

    def _register_metadata_handlers(self) -> Dict[str, Tuple[Callable, Callable]]:
        """Creates a registry mapping file extensions to metadata handlers."""
        return {
            '.pdf': (self._get_pdf_metadata, self._set_pdf_metadata),
            '.jpg': (self._get_image_metadata, self._set_image_metadata),
            '.jpeg': (self._get_image_metadata, self._set_image_metadata),
            '.png': (self._get_image_metadata, self._set_image_metadata),
            '.heic': (self._get_image_metadata, self._set_image_metadata),
            '.tiff': (self._get_image_metadata, self._set_image_metadata),
            '.tif': (self._get_image_metadata, self._set_image_metadata),
            '.docx': (self._get_docx_metadata, self._set_docx_metadata),
            '.pptx': (self._get_pptx_metadata, self._set_pptx_metadata),
            '.xlsx': (self._get_xlsx_metadata, self._set_xlsx_metadata),
        }

    def get_file_metadata(self, file_path: Path) -> Dict[str, Any]:
        """Reads metadata from a file using a registered handler."""
        ext = file_path.suffix.lower()
        handler, _ = self._metadata_handlers.get(ext, (None, None))
        if not handler:
            logger.warning(f"Metadata extraction not supported for {ext} files.")
            return {}

        try:
            return handler(file_path)
        except (DocxPackageNotFoundError, PptxPackageNotFoundError, OpenpyxlInvalidFileException) as e:
            logger.error(f"Cannot open '{file_path.name}'. It might be corrupt or not a valid {ext} file. Error: {e}")
            return {}
        except Exception as e:
            logger.error(f"Failed to read metadata for {file_path.name}: {e}", exc_info=True)
            return {}

    def set_file_metadata(self, file_path: Path, metadata: Dict[str, Any]) -> bool:
        """Writes metadata to a file using a registered handler."""
        ext = file_path.suffix.lower()
        _, handler = self._metadata_handlers.get(ext, (None, None))
        if not handler:
            logger.warning(f"Metadata setting not supported for {ext} files.")
            return False

        try:
            return handler(file_path, metadata)
        except (DocxPackageNotFoundError, PptxPackageNotFoundError, OpenpyxlInvalidFileException) as e:
            logger.error(f"Cannot open '{file_path.name}' to set metadata. Error: {e}")
            return False
        except Exception as e:
            logger.error(f"Failed to set metadata for {file_path.name}: {e}", exc_info=True)
            return False

    # --- Individual Metadata Handlers --- #

    def _safe_decode(self, byte_string: Optional[bytes]) -> str:
        if byte_string is None: return ""
        try: return byte_string.decode('utf-8', errors='replace')
        except AttributeError: return str(byte_string)

    def _get_pdf_metadata(self, file_path: Path) -> Dict[str, Any]:
        with fitz.open(file_path) as doc:
            meta_doc = doc.metadata
            metadata = {
                COMMON_METADATA_KEYS["title"]: meta_doc.get("title", ""),
                COMMON_METADATA_KEYS["author"]: meta_doc.get("author", ""),
                COMMON_METADATA_KEYS["subject"]: meta_doc.get("subject", ""),
                COMMON_METADATA_KEYS["keywords"]: meta_doc.get("keywords", ""),
                COMMON_METADATA_KEYS["creator_tool"]: meta_doc.get("creator", ""),
                COMMON_METADATA_KEYS["creation_date"]: meta_doc.get("creationDate", ""),
                COMMON_METADATA_KEYS["modification_date"]: meta_doc.get("modDate", ""),
            }
            return {k: v for k, v in metadata.items() if v}

    def _set_pdf_metadata(self, file_path: Path, metadata: Dict[str, Any]) -> bool:
        with fitz.open(file_path) as doc:
            update_dict = {
                "title": metadata.get(COMMON_METADATA_KEYS["title"]),
                "author": metadata.get(COMMON_METADATA_KEYS["author"]),
                "subject": metadata.get(COMMON_METADATA_KEYS["subject"]),
                "keywords": metadata.get(COMMON_METADATA_KEYS["keywords"]),
                "creator": metadata.get(COMMON_METADATA_KEYS["creator_tool"]),
            }
            update_dict = {k: v for k, v in update_dict.items() if v is not None}
            if not update_dict:
                return True
            doc.set_metadata(update_dict)
            doc.save(str(file_path), incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
        return True

    def _get_image_metadata(self, file_path: Path) -> Dict[str, Any]:
        metadata = {}
        with Image.open(file_path) as img:
            img_info = img.info
            if file_path.suffix.lower() in (".jpg", ".jpeg", ".heic", ".tiff", ".tif") and "exif" in img_info:
                exif_dict = piexif.load(img_info["exif"])
                if exif_dict.get("0th"):
                    metadata[COMMON_METADATA_KEYS["author"]] = self._safe_decode(exif_dict["0th"].get(piexif.ImageIFD.Artist))
                    metadata[COMMON_METADATA_KEYS["description"]] = self._safe_decode(exif_dict["0th"].get(piexif.ImageIFD.ImageDescription))
                if exif_dict.get("Exif"):
                    metadata[COMMON_METADATA_KEYS["comment"]] = self._safe_decode(exif_dict["Exif"].get(piexif.ExifIFD.UserComment))
            elif file_path.suffix.lower() == ".png":
                for k, v in img_info.items():
                    if isinstance(v, str) and k.lower() in ["author", "title", "description"]:
                        metadata[COMMON_METADATA_KEYS.get(k.lower(), k)] = v
        return {k: v for k, v in metadata.items() if v}

    def _set_image_metadata(self, file_path: Path, metadata: Dict[str, Any]) -> bool:
        with Image.open(file_path) as img:
            img_info = img.info.copy()
            modified = False
            if file_path.suffix.lower() in (".jpg", ".jpeg", ".heic", ".tiff", ".tif"):
                exif_dict = piexif.load(img_info.get("exif", b''))
                if COMMON_METADATA_KEYS["author"] in metadata:
                    exif_dict["0th"][piexif.ImageIFD.Artist] = metadata[COMMON_METADATA_KEYS["author"]].encode()
                    modified = True
                if COMMON_METADATA_KEYS["description"] in metadata:
                    exif_dict["0th"][piexif.ImageIFD.ImageDescription] = metadata[COMMON_METADATA_KEYS["description"]].encode()
                    modified = True
                if COMMON_METADATA_KEYS["comment"] in metadata:
                    exif_dict["Exif"][piexif.ExifIFD.UserComment] = metadata[COMMON_METADATA_KEYS["comment"]].encode()
                    modified = True
                if modified:
                    exif_bytes = piexif.dump(exif_dict)
                    img.save(str(file_path), exif=exif_bytes)
            elif file_path.suffix.lower() == ".png":
                pnginfo = PngImagePlugin.PngInfo()
                if COMMON_METADATA_KEYS["author"] in metadata:
                    pnginfo.add_text("Author", metadata[COMMON_METADATA_KEYS["author"]])
                    modified = True
                if COMMON_METADATA_KEYS["title"] in metadata:
                    pnginfo.add_text("Title", metadata[COMMON_METADATA_KEYS["title"]])
                    modified = True
                if COMMON_METADATA_KEYS["description"] in metadata:
                    pnginfo.add_text("Description", metadata[COMMON_METADATA_KEYS["description"]])
                    modified = True
                if modified:
                    img.save(str(file_path), pnginfo=pnginfo)
        return True

    def _get_office_metadata(self, doc) -> Dict[str, Any]:
        cp = doc.core_properties
        metadata = {
            COMMON_METADATA_KEYS["title"]: cp.title,
            COMMON_METADATA_KEYS["author"]: cp.author,
            COMMON_METADATA_KEYS["subject"]: cp.subject,
            COMMON_METADATA_KEYS["keywords"]: cp.keywords,
            COMMON_METADATA_KEYS["last_modified_by"]: cp.last_modified_by,
            COMMON_METADATA_KEYS["creation_date"]: cp.created.isoformat() if cp.created else None,
            COMMON_METADATA_KEYS["modification_date"]: cp.modified.isoformat() if cp.modified else None,
            COMMON_METADATA_KEYS["description"]: cp.comments,
        }
        return {k: v for k, v in metadata.items() if v is not None and v != ""}

    def _set_office_metadata(self, doc, metadata: Dict[str, Any], file_path: Path) -> bool:
        cp = doc.core_properties
        modified = False
        for key, prop_name in {
            COMMON_METADATA_KEYS["title"]: "title",
            COMMON_METADATA_KEYS["author"]: "author",
            COMMON_METADATA_KEYS["subject"]: "subject",
            COMMON_METADATA_KEYS["keywords"]: "keywords",
            COMMON_METADATA_KEYS["last_modified_by"]: "last_modified_by",
            COMMON_METADATA_KEYS["description"]: "comments",
        }.items():
            if key in metadata:
                setattr(cp, prop_name, metadata[key])
                modified = True
        if modified:
            doc.save(str(file_path))
        return True

    def _get_docx_metadata(self, file_path: Path) -> Dict[str, Any]:
        return self._get_office_metadata(Document(file_path))

    def _set_docx_metadata(self, file_path: Path, metadata: Dict[str, Any]) -> bool:
        doc = Document(file_path)
        return self._set_office_metadata(doc, metadata, file_path)

    def _get_pptx_metadata(self, file_path: Path) -> Dict[str, Any]:
        return self._get_office_metadata(Presentation(file_path))

    def _set_pptx_metadata(self, file_path: Path, metadata: Dict[str, Any]) -> bool:
        prs = Presentation(file_path)
        return self._set_office_metadata(prs, metadata, file_path)

    def _get_xlsx_metadata(self, file_path: Path) -> Dict[str, Any]:
        wb = openpyxl.load_workbook(file_path)
        props = wb.properties
        metadata = {
            COMMON_METADATA_KEYS["title"]: props.title,
            COMMON_METADATA_KEYS["author"]: props.creator,
            COMMON_METADATA_KEYS["subject"]: props.subject,
            COMMON_METADATA_KEYS["keywords"]: props.keywords,
            COMMON_METADATA_KEYS["last_modified_by"]: props.lastModifiedBy,
            COMMON_METADATA_KEYS["creation_date"]: props.created.isoformat() if props.created else None,
            COMMON_METADATA_KEYS["modification_date"]: props.modified.isoformat() if props.modified else None,
            COMMON_METADATA_KEYS["description"]: props.description,
        }
        return {k: v for k, v in metadata.items() if v is not None and v != ""}

    def _set_xlsx_metadata(self, file_path: Path, metadata: Dict[str, Any]) -> bool:
        wb = openpyxl.load_workbook(file_path)
        props = wb.properties
        modified = False

        mapping = {
            COMMON_METADATA_KEYS["title"]: "title",
            COMMON_METADATA_KEYS["author"]: "creator",
            COMMON_METADATA_KEYS["subject"]: "subject",
            COMMON_METADATA_KEYS["keywords"]: "keywords",
            COMMON_METADATA_KEYS["last_modified_by"]: "lastModifiedBy",
            COMMON_METADATA_KEYS["description"]: "description",
        }

        for key, prop_name in mapping.items():
            if key in metadata:
                setattr(props, prop_name, metadata[key])
                modified = True
        
        if modified:
            wb.save(str(file_path))
        return True
